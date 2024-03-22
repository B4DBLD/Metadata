import PyPDF2
import os
from docx import Document
from openpyxl import load_workbook

def docxmeta(filepath):
    docx_file = Document(filepath)
    docx_properties = docx_file.core_properties
    metadata = {
        'author': docx_properties.author,
        'title': docx_properties.title,
        'subject': docx_properties.subject,
        'keywords': docx_properties.keywords,
        'last_modified_by': docx_properties.last_modified_by,
        'created': docx_properties.created,
        'modified': docx_properties.modified,
        'revision': docx_properties.revision,
        'version': docx_properties.version
    }
    return metadata

def xlsxmeta(filepath):
    xlsx_file = load_workbook(filepath)
    return xlsx_file.properties

def pdfmeta(filepath):
    pdf_file = PyPDF2.PdfReader(filepath)
    return pdf_file.metadata


def get_directorypath():
    dir_path = input("Enter the directory path: ")
    return dir_path

def identify_files(dir_path):
    files = os.listdir(dir_path)
    extensions = ['.pdf', '.docx', '.xlsx']
    matched_files = []
    for file in files:
        ext = os.path.splitext(file)[1].lower()
        if ext in extensions:
            matched_files.append(file)
            file_path = os.path.join(dir_path, file)
            if ext == '.docx':
                metadata = docxmeta(file_path)
            elif ext == '.xlsx':
                metadata = xlsxmeta(file_path)
            elif ext == '.pdf':
                metadata = pdfmeta(file_path)
            print(f"Metadata for {file}: {metadata}")



if __name__ == "__main__":
    dirpath = get_directorypath()
    if os.path.isdir(dirpath):
        file_list = identify_files(dirpath)
    else:
        print("The directory does not exist.")
